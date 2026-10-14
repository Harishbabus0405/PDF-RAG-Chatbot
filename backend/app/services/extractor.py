"""Document text extraction engine supporting PDF, DOCX, and TXT formats.

Preserves page numbers (PDF), sections/headings (DOCX, TXT), and provides
robust detection of encrypted, scanned, empty, and malformed files.
"""

from pathlib import Path
from typing import Any, Dict, List, Tuple

import docx
import pymupdf

from backend.app.core.exceptions import (
    DocumentEncryptedException,
    EmptyDocumentException,
    MalformedDocumentException,
    ScannedDocumentException,
    ValidationException,
)
from backend.app.core.logging import get_logger

logger = get_logger("documind.extractor")


def validate_magic_bytes(file_bytes: bytes, extension: str) -> None:
    """Verify that file content header matches the expected format signature."""
    ext = extension.lower()
    if ext == ".pdf":
        if not file_bytes.startswith(b"%PDF-"):
            raise ValidationException(
                message="Invalid PDF header: file content does not match standard PDF signature (%PDF-)",
                details={"expected_signature": "%PDF-", "received_header": file_bytes[:10].decode("ascii", errors="replace")},
            )
    elif ext == ".docx":
        # DOCX is an OpenXML ZIP container starting with PK\x03\x04
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise ValidationException(
                message="Invalid DOCX header: file content does not match ZIP/Office OpenXML signature (PK)",
                details={"expected_signature": "PK\\x03\\x04", "received_header": file_bytes[:10].decode("latin-1", errors="replace")},
            )
    elif ext == ".txt":
        # Check first 512 bytes for null byte which signals a binary executable or file
        if b"\x00" in file_bytes[:512]:
            raise ValidationException(
                message="Invalid plain text file: binary null bytes detected in text stream",
                details={"sample_hex": file_bytes[:16].hex()},
            )


def extract_pdf_content(file_path: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Extract page-by-page text from a PDF file using PyMuPDF.

    Returns:
        Tuple of (pages_list, summary_stats_dict)
    """
    try:
        doc = pymupdf.open(str(file_path))
    except Exception as exc:
        logger.warning(f"PyMuPDF open error on {file_path.name}: {exc}")
        raise MalformedDocumentException(
            message=f"Failed to open PDF document: file structure is malformed or corrupted ({exc})",
            details={"filename": file_path.name, "error": str(exc)},
        )

    # Check for encryption
    if doc.is_encrypted:
        # Check if it can be opened without password
        can_decrypt = False
        try:
            can_decrypt = doc.authenticate("") > 0
        except Exception:
            can_decrypt = False

        if not can_decrypt:
            doc.close()
            raise DocumentEncryptedException(
                message=f"PDF '{file_path.name}' is password-protected or encrypted. Decryption password required.",
                details={"filename": file_path.name, "is_encrypted": True},
            )

    page_count = len(doc)
    if page_count == 0:
        doc.close()
        raise EmptyDocumentException(
            message=f"PDF '{file_path.name}' contains 0 pages.",
            details={"filename": file_path.name, "page_count": 0},
        )

    pages: List[Dict[str, Any]] = []
    total_images = 0
    total_text = ""

    for page_idx in range(page_count):
        page = doc[page_idx]
        page_num = page_idx + 1
        page_text = page.get_text("text").strip()
        images = page.get_images()
        total_images += len(images)

        words = len(page_text.split()) if page_text else 0
        char_count = len(page_text)
        total_text += (" " + page_text)

        pages.append({
            "page_number": page_num,
            "char_count": char_count,
            "word_count": words,
            "text": page_text,
        })

    doc.close()
    clean_total = total_text.strip()

    # Detect scanned or empty PDF
    if not clean_total:
        if total_images > 0:
            raise ScannedDocumentException(
                message=(
                    f"PDF '{file_path.name}' appears to be a scanned document containing "
                    f"{total_images} image(s) with no extractable text layer. Optical Character Recognition (OCR) is required."
                ),
                details={"filename": file_path.name, "image_count": total_images, "page_count": page_count},
            )
        raise EmptyDocumentException(
            message=f"PDF '{file_path.name}' contains no selectable text across its {page_count} page(s).",
            details={"filename": file_path.name, "page_count": page_count},
        )

    summary_stats = {
        "page_count": page_count,
        "character_count": len(clean_total),
        "word_count": len(clean_total.split()),
    }
    return pages, summary_stats


def extract_docx_content(file_path: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Extract section and paragraph text from a Word DOCX file using python-docx.

    Returns:
        Tuple of (sections_list, summary_stats_dict)
    """
    try:
        doc = docx.Document(str(file_path))
    except Exception as exc:
        logger.warning(f"python-docx parse error on {file_path.name}: {exc}")
        raise MalformedDocumentException(
            message=f"Failed to open DOCX document: file is corrupted or not a valid Word archive ({exc})",
            details={"filename": file_path.name, "error": str(exc)},
        )

    sections: List[Dict[str, Any]] = []
    current_heading = "Introduction"
    current_paragraphs: List[str] = []
    all_text_fragments: List[str] = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        # Check for heading style
        style_name = p.style.name if p.style and p.style.name else ""
        if style_name.lower().startswith("heading") or style_name.lower() == "title":
            if current_paragraphs:
                combined_sec = "\n\n".join(current_paragraphs)
                sections.append({
                    "section_index": len(sections) + 1,
                    "heading": current_heading,
                    "char_count": len(combined_sec),
                    "word_count": len(combined_sec.split()),
                    "text": combined_sec,
                })
                current_paragraphs = []
            current_heading = text
        else:
            current_paragraphs.append(text)
            all_text_fragments.append(text)

    # Extract tables content
    for table_idx, table in enumerate(doc.tables):
        table_lines: List[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_lines.append(" | ".join(cells))
        if table_lines:
            tbl_block = f"[Table {table_idx + 1}]\n" + "\n".join(table_lines)
            current_paragraphs.append(tbl_block)
            all_text_fragments.append(tbl_block)

    # Flush final section
    if current_paragraphs:
        combined_sec = "\n\n".join(current_paragraphs)
        sections.append({
            "section_index": len(sections) + 1,
            "heading": current_heading,
            "char_count": len(combined_sec),
            "word_count": len(combined_sec.split()),
            "text": combined_sec,
        })

    total_text = " ".join(all_text_fragments).strip()
    if not total_text:
        raise EmptyDocumentException(
            message=f"DOCX '{file_path.name}' contains no readable paragraph or table text.",
            details={"filename": file_path.name},
        )

    # Estimate page count (approx 500 words or 3000 chars per standard page)
    words = len(total_text.split())
    estimated_pages = max(1, (words + 499) // 500)

    summary_stats = {
        "page_count": estimated_pages,
        "character_count": len(total_text),
        "word_count": words,
    }
    return sections, summary_stats


def extract_txt_content(file_path: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Safely decode and extract text sections from a plain text file.

    Returns:
        Tuple of (sections_list, summary_stats_dict)
    """
    try:
        raw_bytes = file_path.read_bytes()
    except Exception as exc:
        logger.warning(f"Failed to read TXT file {file_path.name}: {exc}")
        raise MalformedDocumentException(
            message=f"Failed to read plain text file ({exc})",
            details={"filename": file_path.name, "error": str(exc)},
        )

    if len(raw_bytes) == 0:
        raise EmptyDocumentException(
            message=f"TXT file '{file_path.name}' is empty (0 bytes).",
            details={"filename": file_path.name, "size_bytes": 0},
        )

    # Check for binary null bytes
    if b"\x00" in raw_bytes[:1024]:
        raise MalformedDocumentException(
            message=f"File '{file_path.name}' contains binary null bytes and cannot be decoded as text.",
            details={"filename": file_path.name},
        )

    # Try encodings in priority order
    encodings = ["utf-8-sig", "utf-8", "latin-1", "cp1252", "iso-8859-1"]
    decoded: str | None = None

    for enc in encodings:
        try:
            decoded = raw_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue

    if decoded is None:
        raise MalformedDocumentException(
            message=f"Unable to decode text file '{file_path.name}' using standard encodings (UTF-8, Latin-1, CP1252).",
            details={"filename": file_path.name},
        )

    # Normalize line endings for cross-platform consistency
    normalized = decoded.replace("\r\n", "\n").replace("\r", "\n")
    clean_text = normalized.strip()
    if not clean_text:
        raise EmptyDocumentException(
            message=f"TXT file '{file_path.name}' contains only whitespace.",
            details={"filename": file_path.name},
        )

    # Group into logical paragraph blocks by double newline or multiple newlines
    raw_blocks = [b.strip() for b in clean_text.split("\n\n") if b.strip()]
    if not raw_blocks:
        raw_blocks = [clean_text]

    sections: List[Dict[str, Any]] = []
    for idx, block in enumerate(raw_blocks):
        sections.append({
            "section_index": idx + 1,
            "heading": f"Section {idx + 1}",
            "char_count": len(block),
            "word_count": len(block.split()),
            "text": block,
        })

    words = len(clean_text.split())
    estimated_pages = max(1, (words + 499) // 500)

    summary_stats = {
        "page_count": estimated_pages,
        "character_count": len(clean_text),
        "word_count": words,
    }
    return sections, summary_stats


def extract_document_text(file_path: Path, extension: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], Dict[str, Any]]:
    """Dispatch document extraction based on extension.

    Returns:
        Tuple of (pages_list, sections_list, summary_stats_dict)
    """
    ext = extension.lower()
    if ext == ".pdf":
        pages, stats = extract_pdf_content(file_path)
        return pages, [], stats
    elif ext == ".docx":
        sections, stats = extract_docx_content(file_path)
        return [], sections, stats
    elif ext == ".txt":
        sections, stats = extract_txt_content(file_path)
        return [], sections, stats
    else:
        raise ValidationException(
            message=f"Unsupported format '{ext}'. Supported formats: .pdf, .docx, .txt",
            details={"extension": ext},
        )
