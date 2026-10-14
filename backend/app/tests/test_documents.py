"""Comprehensive test suite for Phase 4: Document Upload, Extraction, and Lifecycle.

Tests all supported formats (PDF via PyMuPDF, DOCX via python-docx, TXT),
format/size validation, duplicate handling, metadata and text extraction,
error conditions (scanned, encrypted, malformed, empty, oversized),
path traversal protection, and lifecycle operations.
"""

import io
from pathlib import Path
from typing import Generator
import pytest
from fastapi.testclient import TestClient

import docx
import pymupdf

from backend.app.core.config import get_settings


# ==============================================================================
# In-Memory Sample Document Generators
# ==============================================================================

def make_sample_pdf(text: str = "DocuMind AI Automated Intelligence Architecture", pages: int = 1) -> bytes:
    """Create a minimal valid in-memory PDF with text."""
    doc = pymupdf.open()
    for p_idx in range(pages):
        page = doc.new_page()
        page.insert_text((50, 72), f"{text} - Page {p_idx + 1}")
    return doc.tobytes()


def make_encrypted_pdf(password: str = "supersecret123") -> bytes:
    """Create an encrypted PDF protected by a user password."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Confidential financial statement data")
    return doc.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        user_pw=password,
        owner_pw=password,
    )


def make_scanned_pdf() -> bytes:
    """Create a PDF with an image pixmap but no selectable text (scanned document simulation)."""
    doc = pymupdf.open()
    page = doc.new_page()
    # Create a 100x100 RGB pixmap
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 100), 1)
    page.insert_image(pymupdf.Rect(50, 50, 200, 200), pixmap=pix)
    return doc.tobytes()


def make_sample_docx(title: str = "DocuMind System Spec", body: str = "Enterprise neural retrieval engine.") -> bytes:
    """Create a minimal valid DOCX with a heading, paragraph, and table."""
    doc = docx.Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Metric"
    tbl.cell(0, 1).text = "Target"
    tbl.cell(1, 0).text = "Extraction Latency"
    tbl.cell(1, 1).text = "< 50ms"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_sample_txt(content: str = "DocuMind TXT Overview\n\nSection on text extraction pipelines.") -> bytes:
    """Create a valid UTF-8 TXT document."""
    return content.encode("utf-8")


# ==============================================================================
# Document Upload and Extraction Tests
# ==============================================================================

class TestDocumentUploadAndExtraction:
    """Tests covering document ingestion, text extraction, and page/section breakdowns."""

    def test_upload_valid_pdf(self, client: TestClient):
        """Upload a valid multi-page PDF and verify extraction of pages and metadata."""
        pdf_bytes = make_sample_pdf(text="Machine Learning System Architecture", pages=2)
        response = client.post(
            "/api/v1/documents/upload",
            files={"file": ("architecture_spec.pdf", pdf_bytes, "application/pdf")},
        )
        assert response.status_code == 201
        data = response.json()
        assert data["success"] is True
        doc_info = data["data"]
        assert doc_info["id"].startswith("doc_")
        assert doc_info["filename"] == "architecture_spec.pdf"
        assert doc_info["file_type"] == "pdf"
        assert doc_info["status"] == "ready"
        assert doc_info["page_count"] == 2
        assert doc_info["word_count"] > 0
        assert doc_info["character_count"] > 0

    def test_upload_valid_docx(self, client: TestClient):
        """Upload a valid DOCX and verify extraction of sections and metadata."""
        docx_bytes = make_sample_docx(
            title="Knowledge Ingestion Overview",
            body="This document tests Microsoft Word document parsing via python-docx.",
        )
        response = client.post(
            "/api/v1/documents/upload",
            files={"file": ("knowledge_base.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 201
        data = response.json()
        assert data["success"] is True
        doc_info = data["data"]
        assert doc_info["id"].startswith("doc_")
        assert doc_info["filename"] == "knowledge_base.docx"
        assert doc_info["file_type"] == "docx"
        assert doc_info["status"] == "ready"
        assert doc_info["word_count"] > 0

    def test_upload_valid_txt(self, client: TestClient):
        """Upload a valid TXT file and verify paragraph block splitting."""
        txt_bytes = make_sample_txt("DocuMind Notes\n\nSecond paragraph block for safe text ingestion.")
        response = client.post(
            "/api/v1/documents/upload",
            files={"file": ("pipeline_notes.txt", txt_bytes, "text/plain")},
        )
        assert response.status_code == 201
        data = response.json()
        assert data["success"] is True
        doc_info = data["data"]
        assert doc_info["id"].startswith("doc_")
        assert doc_info["filename"] == "pipeline_notes.txt"
        assert doc_info["file_type"] == "txt"
        assert doc_info["status"] == "ready"
        assert doc_info["word_count"] > 0

    def test_duplicate_filenames_receive_unique_ids(self, client: TestClient):
        """Upload two documents with identical original filenames; verify unique IDs and non-collision."""
        pdf_bytes = make_sample_pdf(text="Duplicate collision test payload")
        res1 = client.post(
            "/api/v1/documents/upload",
            files={"file": ("duplicate_report.pdf", pdf_bytes, "application/pdf")},
        )
        res2 = client.post(
            "/api/v1/documents/upload",
            files={"file": ("duplicate_report.pdf", pdf_bytes, "application/pdf")},
        )
        assert res1.status_code == 201
        assert res2.status_code == 201
        id1 = res1.json()["data"]["id"]
        id2 = res2.json()["data"]["id"]
        assert id1 != id2

        # Verify both documents exist simultaneously in the detail endpoint
        detail1 = client.get(f"/api/v1/documents/{id1}").json()
        detail2 = client.get(f"/api/v1/documents/{id2}").json()
        assert detail1["data"]["id"] == id1
        assert detail2["data"]["id"] == id2


# ==============================================================================
# Document Detail and Lifecycle Tests
# ==============================================================================

class TestDocumentLifecycle:
    """Tests covering document retrieval, content verification, pagination, and deletion."""

    def test_get_document_detail_pdf(self, client: TestClient):
        """Verify GET /api/v1/documents/{id} returns full page breakdown with page numbers for PDF."""
        pdf_bytes = make_sample_pdf(text="Page extraction verification test", pages=2)
        upload_res = client.post(
            "/api/v1/documents/upload",
            files={"file": ("paged_test.pdf", pdf_bytes, "application/pdf")},
        )
        doc_id = upload_res.json()["data"]["id"]

        detail_res = client.get(f"/api/v1/documents/{doc_id}")
        assert detail_res.status_code == 200
        data = detail_res.json()["data"]
        assert data["id"] == doc_id
        assert data["page_count"] == 2
        assert len(data["pages"]) == 2
        assert data["pages"][0]["page_number"] == 1
        assert "Page extraction verification test" in data["pages"][0]["text"]
        assert data["pages"][1]["page_number"] == 2

    def test_get_document_detail_docx(self, client: TestClient):
        """Verify GET /api/v1/documents/{id} returns sections with headings for DOCX."""
        docx_bytes = make_sample_docx(title="Microservices Architecture", body="Service-oriented document.")
        upload_res = client.post(
            "/api/v1/documents/upload",
            files={"file": ("microservices.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        doc_id = upload_res.json()["data"]["id"]

        detail_res = client.get(f"/api/v1/documents/{doc_id}")
        assert detail_res.status_code == 200
        data = detail_res.json()["data"]
        assert data["id"] == doc_id
        assert len(data["sections"]) >= 1
        assert data["sections"][0]["heading"] == "Microservices Architecture"

    def test_list_documents_pagination_and_status(self, client: TestClient):
        """Verify GET /api/v1/documents pagination and status filtering."""
        # Ensure at least one ready document exists
        pdf_bytes = make_sample_pdf(text="Pagination test document")
        client.post(
            "/api/v1/documents/upload",
            files={"file": ("pagination_doc.pdf", pdf_bytes, "application/pdf")},
        )

        res = client.get("/api/v1/documents?limit=2&offset=0&status=ready")
        assert res.status_code == 200
        data = res.json()["data"]
        assert data["total"] >= 1
        assert len(data["documents"]) <= 2
        for doc in data["documents"]:
            assert doc["status"] == "ready"

    def test_delete_document_and_verify_404(self, client: TestClient):
        """Upload a document, delete it, and confirm subsequent retrieval returns 404."""
        txt_bytes = make_sample_txt("Temporary document slated for lifecycle deletion.")
        upload_res = client.post(
            "/api/v1/documents/upload",
            files={"file": ("to_delete.txt", txt_bytes, "text/plain")},
        )
        doc_id = upload_res.json()["data"]["id"]

        # Delete document
        delete_res = client.delete(f"/api/v1/documents/{doc_id}")
        assert delete_res.status_code == 200
        del_data = delete_res.json()["data"]
        assert del_data["deleted"] is True
        assert del_data["id"] == doc_id

        # Verify subsequent GET returns 404
        get_res = client.get(f"/api/v1/documents/{doc_id}")
        assert get_res.status_code == 404
        assert get_res.json()["error_code"] == "RESOURCE_NOT_FOUND"

    def test_delete_nonexistent_document_returns_404(self, client: TestClient):
        """Attempting to delete a missing document ID returns 404."""
        res = client.delete("/api/v1/documents/doc_nonexistent_99999")
        assert res.status_code == 404
        assert res.json()["error_code"] == "RESOURCE_NOT_FOUND"

    def test_get_nonexistent_document_returns_404(self, client: TestClient):
        """Attempting to retrieve a missing document ID returns 404."""
        res = client.get("/api/v1/documents/doc_nonexistent_99999")
        assert res.status_code == 404
        assert res.json()["error_code"] == "RESOURCE_NOT_FOUND"


# ==============================================================================
# Error Detection and Format Validation Tests
# ==============================================================================

class TestDocumentValidationErrors:
    """Tests covering error cases: empty files, malformed formats, encrypted PDFs, scanned PDFs."""

    def test_upload_empty_file_rejected(self, client: TestClient):
        """0-byte file must be rejected with 422 validation error."""
        res = client.post(
            "/api/v1/documents/upload",
            files={"file": ("empty.pdf", b"", "application/pdf")},
        )
        assert res.status_code == 422
        data = res.json()
        assert data["success"] is False
        assert "empty" in data["message"].lower()

    def test_upload_unsupported_extension_rejected(self, client: TestClient):
        """Files with unsupported extensions (.exe, .csv, .zip) must be rejected with 422."""
        res = client.post(
            "/api/v1/documents/upload",
            files={"file": ("malicious.exe", b"MZ\x90\x00\x03", "application/x-msdownload")},
        )
        assert res.status_code == 422
        data = res.json()
        assert data["success"] is False
        assert "unsupported" in data["message"].lower()

    def test_upload_magic_bytes_mismatch_rejected(self, client: TestClient):
        """File claiming to be .pdf but containing plain text must be rejected with 422."""
        fake_pdf = b"This is just plain text disguised with a .pdf extension"
        res = client.post(
            "/api/v1/documents/upload",
            files={"file": ("fake_document.pdf", fake_pdf, "application/pdf")},
        )
        assert res.status_code == 422
        data = res.json()
        assert data["success"] is False
        assert "signature" in data["message"].lower() or "header" in data["message"].lower()

    def test_upload_malformed_pdf_rejected(self, client: TestClient):
        """PDF header with corrupted internal structures must be rejected as MALFORMED_DOCUMENT."""
        corrupted_pdf = b"%PDF-1.7\nCorrupted garbage binary \x00\xff\xee\xdd\xcc content without xref table %%EOF"
        res = client.post(
            "/api/v1/documents/upload",
            files={"file": ("corrupted.pdf", corrupted_pdf, "application/pdf")},
        )
        assert res.status_code == 422
        data = res.json()
        assert data["error_code"] in ("MALFORMED_DOCUMENT", "DOCUMENT_PROCESSING_ERROR")

    def test_upload_encrypted_pdf_rejected(self, client: TestClient):
        """Password-protected PDF must be rejected with DOCUMENT_ENCRYPTED error code."""
        encrypted_pdf = make_encrypted_pdf(password="secretpass")
        res = client.post(
            "/api/v1/documents/upload",
            files={"file": ("protected_document.pdf", encrypted_pdf, "application/pdf")},
        )
        assert res.status_code == 422
        data = res.json()
        assert data["error_code"] == "DOCUMENT_ENCRYPTED"
        assert "password" in data["message"].lower() or "encrypted" in data["message"].lower()

    def test_upload_scanned_pdf_rejected(self, client: TestClient):
        """PDF containing scanned image without selectable text layer must be rejected with SCANNED_DOCUMENT."""
        scanned_pdf = make_scanned_pdf()
        res = client.post(
            "/api/v1/documents/upload",
            files={"file": ("scanned_invoice.pdf", scanned_pdf, "application/pdf")},
        )
        assert res.status_code == 422
        data = res.json()
        assert data["error_code"] == "SCANNED_DOCUMENT"
        assert "scanned" in data["message"].lower() or "ocr" in data["message"].lower()

    def test_upload_oversized_file_rejected(self, client: TestClient, monkeypatch):
        """Verify upload endpoint and pre-validation probe reject file sizes exceeding max limit."""
        # 1. Pre-validation probe rejects oversized file at schema boundary
        res_val = client.post(
            "/api/v1/documents/validate",
            json={"filename": "large_archive.pdf", "file_size_bytes": 60 * 1024 * 1024},
        )
        assert res_val.status_code == 422
        assert res_val.json()["error_code"] == "REQUEST_VALIDATION_ERROR"

        # 2. Upload endpoint rejects oversized payload
        from backend.app.core.config import get_settings
        settings = get_settings()
        monkeypatch.setattr(settings, "MAX_UPLOAD_SIZE_BYTES", 500)

        pdf_bytes = make_sample_pdf(text="Exceeding dynamic limit payload")
        res_upload = client.post(
            "/api/v1/documents/upload",
            files={"file": ("too_large.pdf", pdf_bytes, "application/pdf")},
        )
        assert res_upload.status_code == 422
        data = res_upload.json()
        assert data["error_code"] == "VALIDATION_ERROR"
        assert "exceeds maximum limit" in data["message"].lower()


# ==============================================================================
# Security and Path Traversal Tests
# ==============================================================================

class TestPathTraversalAndSecurity:
    """Tests verifying that path traversal and dangerous characters are prevented."""

    def test_path_traversal_in_get_document_id(self, client: TestClient):
        """Attempting path traversal via encoded or raw slashes in document ID returns 400 or 404."""
        # Testing URL encoded traversal sequences
        res = client.get("/api/v1/documents/..%2F..%2Fwindows%2Fsystem.ini")
        assert res.status_code in (400, 404)
        if res.status_code == 400:
            assert res.json()["error_code"] == "INVALID_DOCUMENT_ID"

    def test_path_traversal_in_delete_document_id(self, client: TestClient):
        """Attempting path traversal via DELETE returns 400 or 404."""
        res = client.delete("/api/v1/documents/..%2F..%2Fetc%2Fpasswd")
        assert res.status_code in (400, 404)
        if res.status_code == 400:
            assert res.json()["error_code"] == "INVALID_DOCUMENT_ID"

    def test_malformed_document_id_regex_rejected(self, client: TestClient):
        """Document IDs with illegal characters or too short fail regex validation with 400."""
        res = client.get("/api/v1/documents/short")
        assert res.status_code == 400
        assert res.json()["error_code"] == "INVALID_DOCUMENT_ID"

        res_chars = client.get("/api/v1/documents/doc_!@#$illegalchars")
        assert res_chars.status_code == 400
        assert res_chars.json()["error_code"] == "INVALID_DOCUMENT_ID"

    def test_pre_validate_endpoint_valid_request(self, client: TestClient):
        """POST /api/v1/documents/validate returns success for valid file properties."""
        res = client.post(
            "/api/v1/documents/validate",
            json={"filename": "test_presentation.docx", "file_size_bytes": 1024 * 1024},
        )
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert data["data"]["is_valid"] is True
        assert data["data"]["file_type"] == "docx"
